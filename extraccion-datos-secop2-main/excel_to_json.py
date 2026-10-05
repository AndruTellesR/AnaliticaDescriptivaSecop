import json
import sys
from pathlib import Path

import openpyxl


def excel_to_json(excel_path: str) -> dict:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    definicion = (ws["A1"].value or "").strip()
    api_soda3 = (ws["B1"].value or "").strip()

    columnas = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        num, nombre, descripcion, campo_api, tipo = row[:5]
        if num is None:
            break
        columnas.append({
            "nombreColumna": (nombre or "").strip(),
            "descripcion": (descripcion or "").strip(),
            "nombreCampoApi": (campo_api or "").strip(),
            "tipo": (tipo or "").strip(),
        })

    return {
        "apiSODA3": api_soda3,
        "definicion": definicion,
        "columnas": columnas,
    }


DATASET_PATH = Path(__file__).parent / "dataset.json"


def update_dataset(key: str, excel_path: str):
    with open(DATASET_PATH, "r", encoding="utf-8") as f:
        dataset = json.load(f)

    data = excel_to_json(excel_path)

    existing = dataset.setdefault("fuentes", {}).get(key, {})
    existing.update(data)
    dataset["fuentes"][key] = existing

    with open(DATASET_PATH, "w", encoding="utf-8") as f:
        json.dump(dataset, f, indent=2, ensure_ascii=False)

    print(f'"{key}" actualizado en dataset.json ({len(data["columnas"])} columnas)')


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python excel_to_json.py <key> <ruta_archivo.xlsx>")
        print("Ejemplo: python excel_to_json.py contratosElectronicos definciones/contratos-electronicos.xlsx")
        sys.exit(1)

    key = sys.argv[1]
    path = sys.argv[2]
    update_dataset(key, path)
